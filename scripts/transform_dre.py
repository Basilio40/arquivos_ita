"""
Transform DRE Excel reports (Focco) into dashboard JSON format.
Reads DRE_01_<ano>.xlsx .. DRE_12_<ano>.xlsx and produces:
  - dashboard.json
  - dre-details.json

Optionally overrides the macro revenue value with the sum of the
"Valor Liquido (A Vista) sem o Per" column from the monthly carteira
reports (relatorios_carteira_mensal/Mes_<ano>.xlsx).

CÓPIA VENDORIZADA para o repo de dados (arquivos_ita), a partir de
pro-dash-mp/scripts/transform_dre.py, com três diferenças:
  1. --output-dir  → grava os JSON em pasta arbitrária (default: o
     comportamento original, <project-dir>/src/data).
  2. --ano         → o ano deixa de ser hardcoded nos nomes de arquivo
     DRE_{mm}_{ano}.xlsx / {Mes}_{ano}.xlsx e no regex do rodapé.
  3. MONTH_NUMS cobre os 12 meses (o original ia só até Ago e descartava
     silenciosamente DRE de Set em diante).
"""
import argparse
import openpyxl
import json
import os
import re
from collections import defaultdict

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DEFAULT_SOURCE_DIR = os.path.join(BASE_DIR, 'loja_20274_MP_FIN', 'relatorios_focco')
DEFAULT_PROJECT_DIR = os.path.join(BASE_DIR, 'pro-dash-mp')
DEFAULT_EMPRESA = 'ITALINEA - MP SERV E INTER LTDA'
DEFAULT_PERIODO = 'Janeiro a Julho / 2026'

MONTHS = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul']
MONTH_NUMS = {1: 'Jan', 2: 'Fev', 3: 'Mar', 4: 'Abr', 5: 'Mai', 6: 'Jun', 7: 'Jul',
              8: 'Ago', 9: 'Set', 10: 'Out', 11: 'Nov', 12: 'Dez'}

# Ano dos relatórios (DRE_{mm}_{ANO}.xlsx / {Mes}_{ANO}.xlsx). Definido pelo
# --ano do main(); módulo-global no mesmo estilo do MONTHS/detect_months.
ANO = 2026


def detect_months(source_dir):
    """Meses disponíveis a partir dos DRE_XX_{ANO}.xlsx presentes em source_dir
    (em ordem cronológica). Redefine MONTHS conforme os relatórios existentes.
    """
    global MONTHS
    months = []
    for i in range(1, 13):
        if os.path.exists(os.path.join(source_dir, f'DRE_{i:02d}_{ANO}.xlsx')):
            name = MONTH_NUMS.get(i)
            if name:
                months.append(name)
    if months:
        MONTHS = months
    return months


def parse_brl(val):
    if val is None or val == '':
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s:
        return 0.0
    s = s.replace('R$', '').replace(' ', '').replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def extract_code(acc_str):
    if not acc_str:
        return ''
    s = str(acc_str).strip()
    m = re.match(r'^([\d.]+)', s)
    return m.group(1).rstrip('.') if m else ''


def extract_name(acc_str):
    if not acc_str:
        return ''
    s = str(acc_str).strip()
    m = re.match(r'^[\d.]+\s*(.*)', s)
    return m.group(1).strip() if m else s


def group_code(code):
    if not code:
        return ''
    parts = code.split('.')
    if len(parts) >= 2:
        return f'{parts[0]}.{parts[1]}'
    return parts[0] if parts else ''


def group_top(code):
    if not code:
        return ''
    if code.startswith('3'):
        return 'RECEITAS'
    elif code.startswith('4'):
        return 'DESPESAS'
    return ''


def project_key(cc, titulo):
    """Número do projeto (chave de cliente) — do Centro de Custo se houver,
    senão dos dígitos iniciais do "Título/Parcela".

    O Focco grava o CC como "<nº do projeto><nome>" (ex.: "56801ArytanaDAST");
    quando o CC vem vazio, o mesmo número aparece no início do título
    ("56801/0"). Aceita prefixo de 4+ dígitos para ser robusto a variações.
    """
    for src in (cc, titulo):
        if not src:
            continue
        m = re.match(r'^(\d{4,})', str(src).strip())
        if m:
            return m.group(1)
    return ''


def client_label(key, cc, parte):
    """Rótulo legível do cliente para o detalhamento.

    Prioriza o Centro de Custo original (já vem com o nome do sistema Focco).
    Quando o CC está vazio, monta "<nº> <Cliente/Fornecedor>" para os projetos
    que só aparecem sem CC; se não houver 'parte', devolve só o número.
    """
    if cc:
        return str(cc).strip()
    if key and parte:
        return f'{key} {str(parte).strip()}'
    return key


# Grupos de despesa considerados no detalhamento por cliente e seus rótulos.
EXPENSE_GROUPS = ('4.02', '4.03', '4.07')
CATEGORY_LABELS = {
    '4.02': 'Compra e frete',
    '4.03': 'Comissões',
    '4.07': 'Financeiro',
}


def expense_category(code):
    """Categoria de exibição do detalhamento a partir do código de conta.

    'Receita' para contas 3.x; para despesas, mapeia o grupo (4.02/4.03/4.07)
    para um rótulo legível. Retorna '' se o grupo não for relevante.
    """
    if not code:
        return ''
    if code.startswith('3.'):
        return 'Receita'
    grupo = '.'.join(code.split('.')[:2])
    return CATEGORY_LABELS.get(grupo, '')


def month_from_filename(fname):
    m = re.search(rf'DRE_(\d{{2}})_{ANO}', fname)
    if m:
        return MONTH_NUMS.get(int(m.group(1)), '')
    return ''


def is_summary_row(row):
    acc = row[0]
    total = row[1]
    titulo = row[4] if len(row) > 4 else None
    return bool(acc) and bool(total) and not titulo


def is_detail_row(row):
    titulo = row[4] if len(row) > 4 else None
    return bool(titulo)


def process_file(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    month = month_from_filename(os.path.basename(filepath))
    summaries = {}
    detail_rows = []
    break_even = 0.0

    for row in ws.iter_rows(min_row=17, values_only=True):
        if not row or not row[0]:
            continue
        acc_str = str(row[0]).strip()
        if not acc_str:
            continue
        code = extract_code(acc_str)
        if code == '7.70' and is_summary_row(row):
            break_even = parse_brl(row[1])
            continue
        if code.startswith('7'):
            continue
        if is_summary_row(row):
            total = parse_brl(row[1])
            summaries[code] = total
            continue
        if is_detail_row(row):
            valor = parse_brl(row[8])
            if valor == 0:
                valor = parse_brl(row[15]) if len(row) > 15 else 0
            data_mov = str(row[9]) if len(row) > 9 and row[9] else ''
            detail_rows.append({
                'code': code,
                'group': group_code(code),
                'top': group_top(code),
                'acc_name': extract_name(acc_str),
                'titulo': str(row[4] or ''),
                'tip_doc': str(row[5] or ''),
                'portador': str(row[6] or ''),
                'parte': str(row[7] or ''),
                'valor': valor,
                'data_mov': data_mov,
                'obs': str(row[11] or '') if len(row) > 11 else '',
                'cc': str(row[14] or '').strip() if len(row) > 14 else '',
            })

    return month, summaries, detail_rows, break_even


def extract_receita_carteira(filepath):
    """Sum the 'Valor Liquido (A Vista) sem o Per' column (12, 0-based index 11)
    from the detail rows of a monthly carteira report (Mes_2026.xlsx).

    Returns (total, reported_total) where reported_total is the value on the
    'Total Liquido sem o Per' footer row, for divergence reporting.
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    total = 0.0
    reported_total = None
    for row in ws.iter_rows(min_row=31, values_only=True):
        if not row:
            continue
        first = row[0]
        if isinstance(first, str) and first.strip().lower().startswith('total'):
            if 'sem o per' in first.strip().lower():
                reported_total = parse_brl(row[1]) if len(row) > 1 else None
            continue
        valor = row[11] if len(row) > 11 else None
        if isinstance(valor, (int, float)):
            total += float(valor)
    return total, reported_total


def load_carteira_receitas(carteira_dir):
    """Read all Mes_{ANO}.xlsx files in carteira_dir and return {mes: receita}."""
    receitas = {}
    if not carteira_dir or not os.path.isdir(carteira_dir):
        return receitas
    for m in MONTHS:
        fpath = os.path.join(carteira_dir, f'{m}_{ANO}.xlsx')
        if not os.path.exists(fpath):
            print(f'  WARNING: carteira {m}_{ANO}.xlsx not found, skipping')
            continue
        total, reported = extract_receita_carteira(fpath)
        if reported is not None and abs(reported - total) > 0.01:
            print(f'  NOTE: {m}_{ANO}.xlsx detail sum ({total:,.2f}) differs from '
                  f'reported total ({reported:,.2f}); using detail sum')
        receitas[m] = total
        print(f'  Carteira {m}: {total:,.2f}')
    return receitas


GROUP_DEFINITIONS = {
    '3.01': ('RECEITAS', 'RECEITA BRUTA'),
    '4.01': ('DESPESAS', 'IMPOSTOS SOBRE VENDAS'),
    '4.02': ('DESPESAS', 'CUSTO DAS MERCADORIAS VENDIDAS'),
    '4.03': ('DESPESAS', 'DESPESAS COM PESSOAL / COMISSOES'),
    '4.04': ('DESPESAS', 'DESPESAS OPERACIONAIS'),
    '4.07': ('DESPESAS', 'DESPESAS FINANCEIRAS'),
}


def build_dashboard(source_dir, carteira_dir=None,
                    empresa=DEFAULT_EMPRESA, periodo=DEFAULT_PERIODO):
    months = detect_months(source_dir)
    print(f'  Months detected: {months or MONTHS}')
    monthly_summaries = {}
    monthly_details = {}
    monthly_break_even = {}
    all_details = []

    for i in range(1, 13):
        fname = f'DRE_{i:02d}_{ANO}.xlsx'
        fpath = os.path.join(source_dir, fname)
        if not os.path.exists(fpath):
            print(f'  WARNING: {fname} not found, skipping')
            continue
        print(f'  Processing {fname}...')
        month, summaries, details, break_even = process_file(fpath)
        monthly_summaries[month] = summaries
        monthly_details[month] = details
        monthly_break_even[month] = break_even
        all_details.extend([(month, d) for d in details])

    # carteira mensal (optional revenue override)
    carteira_receitas = load_carteira_receitas(carteira_dir) if carteira_dir else {}
    use_carteira = bool(carteira_receitas)
    if use_carteira:
        print('  Using carteira mensal values as revenue')

    # dreMonth
    dre_month = []
    for m in MONTHS:
        sums = monthly_summaries.get(m, {})
        receitas = sum(v for c, v in sums.items() if c.startswith('3.') and c.count('.') == 1)
        if use_carteira and m in carteira_receitas:
            receitas = carteira_receitas[m]
        despesas = sum(v for c, v in sums.items() if c.startswith('4.') and c.count('.') == 1)
        resultado = receitas - despesas
        margem = (resultado / receitas * 100) if receitas else 0
        dre_month.append({
            'mes': m, 'receitas': round(receitas, 2),
            'despesas': round(despesas, 2), 'resultado': round(resultado, 2),
            'margem': round(margem, 1),
            'breakEven': round(monthly_break_even.get(m, 0.0), 2),
        })

    # dreGroups
    all_groups = set()
    for sums in monthly_summaries.values():
        for code in sums:
            gc = group_code(code)
            if gc and (gc.startswith('3.') or gc.startswith('4.')):
                all_groups.add(gc)

    dre_groups = []
    for gc in sorted(all_groups):
        if gc not in GROUP_DEFINITIONS:
            continue
        top, nome = GROUP_DEFINITIONS[gc]
        row_data = {'top': top, 'grupoCod': gc, 'grupoNome': nome}
        total = 0.0
        for m in MONTHS:
            sums = monthly_summaries.get(m, {})
            val = sums.get(gc, 0.0)
            row_data[m] = round(val, 2)
            total += val
        row_data['total'] = round(total, 2)
        dre_groups.append(row_data)
    dre_groups.sort(key=lambda g: (0 if g['top'] == 'RECEITAS' else 1, g['grupoCod']))

    # dre-details
    item_map = {}
    for m in MONTHS:
        sums = monthly_summaries.get(m, {})
        for code, val in sums.items():
            gc = group_code(code)
            if gc not in GROUP_DEFINITIONS:
                continue
            if code.count('.') < 2:
                continue
            key = (gc, code)
            if key not in item_map:
                name = code
                for md, details in monthly_details.items():
                    for d in details:
                        if d['code'] == code:
                            name = d['acc_name']
                            break
                    if name != code:
                        break
                item_map[key] = {'grupoCod': gc, 'itemCod': code, 'itemNome': name}
            if m not in item_map[key]:
                item_map[key][m] = 0.0
            item_map[key][m] += val

    dre_details = []
    for (gc, code), info in sorted(item_map.items()):
        total = sum(info.get(m, 0) for m in MONTHS)
        detail = {'grupoCod': gc, 'itemCod': code, 'itemNome': info['itemNome']}
        for m in MONTHS:
            detail[m] = round(info.get(m, 0), 2)
        detail['total'] = round(total, 2)
        dre_details.append(detail)

    # fluxoCaixa
    fluxo = []
    saldo_acum = 0.0
    for dm in dre_month:
        liquido = dm['resultado']
        saldo_acum += liquido
        fluxo.append({
            'mes': dm['mes'], 'entradasEfet': dm['receitas'], 'entradasPrev': 0,
            'saidasEfet': -dm['despesas'], 'saidasPrev': 0,
            'liquido': round(liquido, 2), 'saldoAcum': round(saldo_acum, 2),
        })

    # carteira
    carteira = []
    for m in MONTHS:
        details = monthly_details.get(m, [])
        rev_rows = [d for d in details if d['top'] == 'RECEITAS']
        clients_in_month = {}
        for d in rev_rows:
            parte = d['parte'].strip()
            if parte:
                clients_in_month[parte] = clients_in_month.get(parte, 0) + d['valor']
        if use_carteira and m in carteira_receitas:
            total = carteira_receitas[m]
        else:
            total = sum(d['valor'] for d in rev_rows)
        num_clients = len(clients_in_month)
        ticket = total / num_clients if num_clients else 0
        carteira.append({
            'mes': m, 'total': round(total, 2), 'novosClientes': num_clients,
            'novosValor': round(total, 2), 'clientesAtivos': num_clients,
            'ticketMedio': round(ticket, 2),
        })

    # clientes
    client_map = {}
    client_first = {}
    for m_idx, m in enumerate(MONTHS):
        details = monthly_details.get(m, [])
        for d in details:
            if d['top'] != 'RECEITAS':
                continue
            parte = d['parte'].strip()
            if not parte:
                continue
            if parte not in client_map:
                client_map[parte] = {mm: 0 for mm in MONTHS}
                client_first[parte] = m
            client_map[parte][m] += d['valor']

    clientes = []
    for name, months_data in client_map.items():
        total = sum(months_data.values())
        if total < 1:
            continue
        row = {'cliente': name, 'total': round(total, 2), 'primeiroMes': client_first[name]}
        for m in MONTHS:
            row[m] = round(months_data[m], 2)
        clientes.append(row)
    clientes.sort(key=lambda c: c['total'], reverse=True)

    # Cost center analysis
    cc_monthly = defaultdict(lambda: {m: 0 for m in MONTHS})
    for m in MONTHS:
        details = monthly_details.get(m, [])
        for d in details:
            if d['top'] != 'DESPESAS':
                continue
            cc = d['cc'] if d['cc'] else 'SEM CC'
            cc_monthly[cc][m] += d['valor']

    cc_totals = {cc: sum(md.values()) for cc, md in cc_monthly.items()}
    sorted_ccs = sorted(cc_totals.items(), key=lambda x: x[1], reverse=True)
    top_cc_names = [cc for cc, _ in sorted_ccs[:8]]

    cc_all = []
    for cc, total in sorted_ccs:
        row = {'centro': cc, 'total': round(total, 2)}
        for m in MONTHS:
            row[m] = round(cc_monthly[cc][m], 2)
        cc_all.append(row)

    cc_top = cc_all[:20]

    cc_stack = []
    other_ccs = [cc for cc, _ in sorted_ccs[8:]]
    for m in MONTHS:
        row = {'mes': m}
        for cc in top_cc_names:
            row[cc] = round(cc_monthly[cc][m], 2)
        row['Outros'] = round(sum(cc_monthly[cc][m] for cc in other_ccs), 2)
        cc_stack.append(row)

    # naoEntregues
    # Agrupa por NÚMERO DO PROJETO (chave de cliente), não pela string crua do
    # Centro de Custo. Assim, quando o CC vem vazio, o cliente é derivado dos
    # dígitos iniciais do "Título/Parcela" e ainda casa despesas (4.02/4.03/4.07)
    # com as receitas (3.01) do mesmo projeto. O rótulo exibido prioriza o CC
    # original (já vem com o nome do sistema Focco); para projetos só vistos sem
    # CC, monta "<nº> <Cliente/Fornecedor>".
    #
    # "A entregar" = receita − total de despesas do projeto (resultado líquido).
    # Projetos quitados (sem receita ou com receita <= despesa) ficam com
    # aEntregar <= 0 e são filtrados pelo front-end (.filter(aEntregar > 0)).
    cc_purchases = defaultdict(list)
    cc_revenue = defaultdict(float)
    cc_totals_cat = defaultdict(lambda: defaultdict(float))  # key -> {categoria: total}
    cc_label = {}
    for m in MONTHS:
        details = monthly_details.get(m, [])
        for d in details:
            is_expense = d['group'] in EXPENSE_GROUPS
            is_revenue = d['top'] == 'RECEITAS' and d['group'] == '3.01'
            if not (is_expense or is_revenue):
                continue
            key = project_key(d['cc'], d['titulo'])
            if not key:
                continue  # sem nº de projeto -> não pertence a um cliente
            if is_expense:
                cc_purchases[key].append((m, d))
                cc_totals_cat[key][expense_category(d['code'])] += d['valor']
            else:
                cc_revenue[key] += d['valor']
            # Rótulo: CC preenchido tem precedência (nome do Focco); fallback
            # só preenche se ainda não houver rótulo melhor.
            if d['cc']:
                cc_label.setdefault(key, str(d['cc']).strip())
            else:
                cc_label.setdefault(key, client_label(key, d['cc'], d['parte']))

    nao_entregues = []
    total_compras = 0      # total de despesas (4.02+4.03+4.07)
    total_a_entregar = 0   # soma de (receita − despesa) por projeto
    num_clientes_ne = 0
    for key, items in cc_purchases.items():
        total_despesa = sum(d['valor'] for _, d in items)
        receita = cc_revenue.get(key, 0.0)
        a_entregar = receita - total_despesa
        if total_despesa < 1:
            continue
        # subtotal por categoria, para o front-end agrupar os itens
        categorias = {cat: round(v, 2) for cat, v in cc_totals_cat[key].items()}
        itens = []
        for m, d in items:
            cat = expense_category(d['code'])
            itens.append({
                'mes': m, 'titulo': d['titulo'], 'obs': d['obs'],
                'valor': round(d['valor'], 2), 'conta': d['acc_name'],
                'fornecedor': d['parte'], 'tipo': cat or 'Compra',
                'categoria': cat, 'tipDoc': d.get('tip_doc', ''),
            })
        nao_entregues.append({
            'centroCusto': cc_label.get(key, key),
            'totalCompra': round(total_despesa, 2),
            'receita': round(receita, 2), 'aEntregar': round(a_entregar, 2),
            'categorias': categorias,
            'itens': itens,
        })
        total_compras += total_despesa
        total_a_entregar += a_entregar
        num_clientes_ne += 1
    # Ordena por "a entregar" desc — as maiores pendências líquidas primeiro
    nao_entregues.sort(key=lambda x: x['aEntregar'], reverse=True)

    # meta
    meta = {
        'empresa': empresa,
        'periodo': periodo,
        'totalRegistros': len(all_details),
    }

    total_receitas = sum(dm['receitas'] for dm in dre_month)
    # O resumo reflete apenas clientes pendentes (aEntregar > 0), para casar
    # com o que o front-end exibe após .filter(aEntregar > 0).
    pendentes = [n for n in nao_entregues if n['aEntregar'] > 0]
    nao_entregues_resumo = {
        'totalCompras': round(sum(n['totalCompra'] for n in pendentes), 2),
        'totalReceitas': round(sum(n['receita'] for n in pendentes), 2),
        'totalAEntregar': round(sum(n['aEntregar'] for n in pendentes), 2),
        'numClientes': len(pendentes),
        'numClientesComDespesa': num_clientes_ne,
    }

    dashboard = {
        'dreMonth': dre_month, 'dreGroups': dre_groups, 'fluxoCaixa': fluxo,
        'naoEntregues': nao_entregues, 'ccStack': cc_stack, 'ccTop': cc_top,
        'ccAll': cc_all, 'carteira': carteira, 'clientes': clientes,
        'topCC': top_cc_names, 'meta': meta, 'naoEntreguesResumo': nao_entregues_resumo,
    }
    return dashboard, dre_details


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description='Transform DRE Excel reports into dashboard JSON')
    parser.add_argument('--source-dir', default=DEFAULT_SOURCE_DIR,
                        help='Directory with DRE_01_<ano>.xlsx .. DRE_12_<ano>.xlsx')
    parser.add_argument('--carteira-dir', default=None,
                        help='Directory with carteira mensal reports (Mes_<ano>.xlsx). '
                             'When provided, its "Valor Liquido (A Vista) sem o Per" '
                             'totals replace the DRE revenue.')
    parser.add_argument('--project-dir', default=DEFAULT_PROJECT_DIR,
                        help='Dashboard project root (writes src/data/*.json when '
                             '--output-dir is not given)')
    parser.add_argument('--output-dir', default=None,
                        help='Arbitrary output directory for the JSONs (overrides '
                             '<project-dir>/src/data)')
    parser.add_argument('--ano', type=int, default=2026,
                        help='Ano dos relatórios (nomes DRE_{mm}_{ano} / {Mes}_{ano})')
    parser.add_argument('--empresa', default=DEFAULT_EMPRESA,
                        help='Company name written to meta.empresa')
    parser.add_argument('--periodo', default=DEFAULT_PERIODO,
                        help='Period label written to meta.periodo')
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    global ANO
    ANO = args.ano
    print('=== DRE Excel to Dashboard JSON Transform ===')
    print(f'  source:   {args.source_dir}')
    print(f'  carteira: {args.carteira_dir or "(none - using DRE revenue)"}')
    print(f'  ano:      {ANO}')
    output_dir = args.output_dir or os.path.join(args.project_dir, 'src', 'data')
    os.makedirs(output_dir, exist_ok=True)
    print(f'  output:   {output_dir}')
    dashboard, dre_details = build_dashboard(
        source_dir=args.source_dir, carteira_dir=args.carteira_dir,
        empresa=args.empresa, periodo=args.periodo)
    dash_path = os.path.join(output_dir, 'dashboard.json')
    with open(dash_path, 'w', encoding='utf-8') as f:
        json.dump(dashboard, f, ensure_ascii=False, separators=(',', ':'))
    print(f'\nWritten: {dash_path}')
    print(f'  dreMonth: {len(dashboard["dreMonth"])} months')
    print(f'  dreGroups: {len(dashboard["dreGroups"])} groups')
    print(f'  fluxoCaixa: {len(dashboard["fluxoCaixa"])} months')
    print(f'  naoEntregues: {len(dashboard["naoEntregues"])} cost centers')
    print(f'  ccStack: {len(dashboard["ccStack"])} rows')
    print(f'  ccTop: {len(dashboard["ccTop"])} centers')
    print(f'  ccAll: {len(dashboard["ccAll"])} centers')
    print(f'  carteira: {len(dashboard["carteira"])} months')
    print(f'  clientes: {len(dashboard["clientes"])} clients')
    print(f'  topCC: {dashboard["topCC"]}')
    print(f'  meta: {dashboard["meta"]}')
    details_path = os.path.join(output_dir, 'dre-details.json')
    with open(details_path, 'w', encoding='utf-8') as f:
        json.dump(dre_details, f, ensure_ascii=False, indent=2)
    print(f'\nWritten: {details_path}')
    print(f'  {len(dre_details)} detail items')


if __name__ == '__main__':
    main()